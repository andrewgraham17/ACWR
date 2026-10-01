"""Storage. Every function writes immediately, so the app autosaves on each change.

Where the data goes:
  * DATABASE_URL set (Streamlit secrets or an environment variable) -> that online database
    (e.g. a free Neon / Supabase Postgres). Use this when the app is deployed on the web.
  * otherwise -> data/throwing.db, a SQLite file next to the app (fine on your own computer).
"""
from __future__ import annotations

import os
import re
import time
import unicodedata
import uuid
from datetime import date, datetime

import pandas as pd
from sqlalchemy import create_engine, text

DB_PATH = os.environ.get("THROW_DB", os.path.join(os.path.dirname(os.path.abspath(__file__)), "data", "throwing.db"))

LOG_FIELDS = ["date", "session_type", "distance", "start_time", "end_time", "actual_throws",
              "velo_low", "velo_high", "warm_throws", "warm_velo_low", "warm_velo_high",
              "sixty_throws", "sixty_velo_low", "sixty_velo_high", "notes"]
NUM_FIELDS = {"distance", "actual_throws", "velo_low", "velo_high", "warm_throws", "warm_velo_low",
              "warm_velo_high", "sixty_throws", "sixty_velo_low", "sixty_velo_high"}

# 2026 pitching staff (jersey, name, position, year, bats/throws)
DEFAULT_ROSTER = [
    ("30", "Max Bettis", "RHP", "So.", "R/R"), ("33", "Sam Bettis", "RHP", "So.", "R/R"),
    ("26", "Noah Gartner", "RHP", "So.", "R/R"), ("31", "Cooper Jesperson", "RHP", "So.", "R/R"),
    ("23", "Jonah Reich", "RHP", "5th", "R/R"), ("32", "Jackson Baker", "RHP", "Sr.", "R/R"),
    ("7", "Trenton Buckley", "RHP", "Jr.", "R/R"), ("17", "Stevie Doty", "RHP", "Jr.", "R/R"),
    ("99", "Colin Driffill", "RHP", "Fr.", "R/R"), ("37", "Jackson Hulcher", "RHP", "Sr.", "R/R"),
    ("35", "Griffin Lewis", "RHP", "Jr.", "R/R"), ("50", "Max Nantais-Vlahovich", "RHP", "Fr.", "R/R"),
    ("28", "Joseph Pereira", "RHP", "Fr.", "L/R"), ("18", "Stocton Timbrook", "RHP", "Jr.", "R/R"),
    ("15", "Grady Westphal", "RHP", "So.", "R/R"), ("4", "Cordell Clinkingbeard", "RHP/INF", "Fr.", "R/R"),
    ("16", "Kyler Horsman", "RHP/OF", "Sr.", "R/R"), ("9", "Donte Lewis", "UTL", "Jr.", "R/R"),
    ("41", "Adam Arther", "LHP", "Sr.", "L/L"), ("0", "Kaden Dean", "LHP", "Fr.", "L/L"),
    ("25", "Grant Meert", "LHP", "Sr.", "L/L"),
]


def slugify(name: str) -> str:
    s = unicodedata.normalize("NFKD", name).lower()
    s = re.sub(r"[^\w\s-]", "", s).strip()
    return re.sub(r"-+", "-", re.sub(r"[\s_]+", "-", s))


def database_url() -> str:
    url = os.environ.get("DATABASE_URL")
    if not url:
        try:
            import streamlit as st
            url = st.secrets.get("DATABASE_URL")
        except Exception:
            url = None
    if url:
        url = url.strip().strip('"')
        if url.startswith("postgres://"):
            url = "postgresql://" + url[len("postgres://"):]
        if url.startswith("postgresql://"):            # use the psycopg2 driver in requirements.txt
            url = "postgresql+psycopg2://" + url[len("postgresql://"):]
        return url
    os.makedirs(os.path.dirname(DB_PATH), exist_ok=True)
    return "sqlite:///" + DB_PATH


def using_online_db() -> bool:
    return not database_url().startswith("sqlite")


_ENGINE = None


def _engine():
    global _ENGINE
    if _ENGINE is None:
        _ENGINE = create_engine(database_url(), pool_pre_ping=True, future=True)
        num = "DOUBLE PRECISION"
        cols = ", ".join(f"{f} {num if f in NUM_FIELDS else 'TEXT'}" for f in LOG_FIELDS)
        with _ENGINE.begin() as c:
            c.execute(text(f"""CREATE TABLE IF NOT EXISTS players(
                id TEXT PRIMARY KEY, name TEXT, number TEXT, position TEXT, year TEXT, bt TEXT,
                avg_top_velo {num}, peak_velo {num}, active INTEGER DEFAULT 1, sort INTEGER,
                target_type TEXT DEFAULT 'Game', target_date TEXT, planned_pitches {num}, planned_innings {num})"""))
            c.execute(text(f"CREATE TABLE IF NOT EXISTS throws(id TEXT PRIMARY KEY, player TEXT, created {num}, updated {num}, {cols})"))
            c.execute(text("CREATE TABLE IF NOT EXISTS overrides(player TEXT, day TEXT, day_type TEXT, PRIMARY KEY(player, day))"))
    return _ENGINE


def _run(sql, **params):
    with _engine().begin() as c:
        return c.execute(text(sql), params)


def _one(sql, **params):
    with _engine().begin() as c:
        r = c.execute(text(sql), params).fetchone()
    return r


def _df(sql, **params) -> pd.DataFrame:
    with _engine().connect() as c:
        return pd.read_sql(text(sql), c, params=params)


def _clean(field, v):
    if v is None or (isinstance(v, float) and pd.isna(v)):
        return None
    if field == "date":
        if isinstance(v, datetime):
            v = v.date()
        return v.isoformat() if isinstance(v, date) else (str(v)[:10] or None)
    if field in ("start_time", "end_time"):
        return v.strftime("%H:%M") if hasattr(v, "strftime") else (str(v).strip()[:5] or None)
    if field in NUM_FIELDS:
        try:
            return float(v)
        except (TypeError, ValueError):
            return None
    s = str(v).strip()
    return s or None


# ───────── players ─────────
def load_players(include_inactive=False) -> pd.DataFrame:
    df = _df("SELECT * FROM players" + ("" if include_inactive else " WHERE active=1") + " ORDER BY sort, name")
    df["target_date"] = [date.fromisoformat(v) if isinstance(v, str) and v else None for v in df["target_date"]]
    df["target_date"] = df["target_date"].astype(object)
    return df


def get_player(pid) -> dict | None:
    df = load_players(include_inactive=True)
    hit = df[df["id"] == pid]
    if hit.empty:
        return None
    r = hit.iloc[0].to_dict()
    for k in ("avg_top_velo", "peak_velo", "planned_pitches", "planned_innings"):
        if pd.isna(r.get(k)):
            r[k] = None
    if not isinstance(r.get("target_date"), date) or pd.isna(r.get("target_date")):
        r["target_date"] = None
    return r


def add_player(name, number="", position="RHP", year=None, bt=None, avg_top=None, peak=None) -> str:
    pid = slugify(name)
    n = _one("SELECT COALESCE(MAX(sort), 0) + 1 FROM players")[0]
    if _one("SELECT 1 FROM players WHERE id=:id", id=pid):   # re-adding a removed player brings his log back
        _run("UPDATE players SET active=1, name=:n, number=:num, position=:pos WHERE id=:id",
             n=name.strip(), num=number, pos=position, id=pid)
        if avg_top is not None:
            _run("UPDATE players SET avg_top_velo=:v WHERE id=:id", v=avg_top, id=pid)
    else:
        _run("""INSERT INTO players(id, name, number, position, year, bt, avg_top_velo, peak_velo, active, sort)
                VALUES (:id, :n, :num, :pos, :yr, :bt, :avg, :peak, 1, :sort)""",
             id=pid, n=name.strip(), num=number, pos=position, yr=year, bt=bt, avg=avg_top, peak=peak, sort=int(n))
    return pid


def update_player(pid, **fields):
    allowed = {"name", "number", "position", "year", "bt", "avg_top_velo", "peak_velo", "active",
               "target_type", "target_date", "planned_pitches", "planned_innings"}
    for k, v in fields.items():
        if k not in allowed:
            continue
        if isinstance(v, date):
            v = v.isoformat()
        if isinstance(v, float) and pd.isna(v):
            v = None
        if hasattr(v, "item"):          # numpy number -> plain Python
            v = v.item()
        if k in ("avg_top_velo", "peak_velo", "planned_pitches", "planned_innings") and v is not None:
            v = float(v)
        if k in ("active", "sort") and v is not None:
            v = int(v)
        _run(f"UPDATE players SET {k}=:v WHERE id=:id", v=v, id=pid)


def remove_player(pid):
    """Hides his tab. His throw log is kept and comes back if he's added again."""
    update_player(pid, active=0)


# ───────── throw log (row level) ─────────
def load_log(pid) -> pd.DataFrame:
    df = _df(f"SELECT id, created, {', '.join(LOG_FIELDS)} FROM throws WHERE player=:p", p=pid)
    df = df.sort_values(["date", "created"], ascending=False, na_position="last").reset_index(drop=True)
    return df


def get_row(row_id) -> dict | None:
    r = _one("SELECT * FROM throws WHERE id=:id", id=row_id)
    return dict(r._mapping) if r else None


def add_row(pid, rec: dict) -> str:
    rid = uuid.uuid4().hex[:12]
    now = time.time()
    vals = {f: _clean(f, rec.get(f)) for f in LOG_FIELDS}
    _run(f"INSERT INTO throws(id, player, created, updated, {', '.join(LOG_FIELDS)}) "
         f"VALUES (:id, :player, :created, :updated, {', '.join(':' + f for f in LOG_FIELDS)})",
         id=rid, player=pid, created=float(rec.get("created") or now), updated=now, **vals)
    return rid


def update_row(row_id, rec: dict):
    fields = [f for f in LOG_FIELDS if f in rec]
    if not fields:
        return
    _run(f"UPDATE throws SET {', '.join(f'{f}=:{f}' for f in fields)}, updated=:updated WHERE id=:id",
         updated=time.time(), id=row_id, **{f: _clean(f, rec[f]) for f in fields})


def delete_row(row_id):
    _run("DELETE FROM throws WHERE id=:id", id=row_id)


# ───────── overrides ─────────
def load_overrides(pid) -> dict:
    with _engine().connect() as c:
        rows = c.execute(text("SELECT day, day_type FROM overrides WHERE player=:p"), {"p": pid}).fetchall()
    return {date.fromisoformat(d): t for d, t in rows}


def save_override(pid, day: date, day_type):
    if not day_type or day_type == "(none)":
        _run("DELETE FROM overrides WHERE player=:p AND day=:d", p=pid, d=day.isoformat())
    else:
        _run("""INSERT INTO overrides(player, day, day_type) VALUES (:p, :d, :t)
                ON CONFLICT(player, day) DO UPDATE SET day_type=excluded.day_type""", p=pid, d=day.isoformat(), t=day_type)


def clear_overrides(pid):
    _run("DELETE FROM overrides WHERE player=:p", p=pid)


# ───────── first run / import ─────────
def seed_if_empty(workbook_path=None):
    if _one("SELECT COUNT(*) FROM players")[0]:
        return False
    for num, name, pos, yr, bt in DEFAULT_ROSTER:
        add_player(name, num, pos, yr, bt)
    if workbook_path and os.path.exists(workbook_path):
        import_workbook(workbook_path)
    return True


def import_workbook(src) -> tuple[int, int]:
    """Import Avg Top Velo / Peak Velo (Roster tab), the Throw Log and planner inputs from the ACWR workbook."""
    import warnings
    import openpyxl
    with warnings.catch_warnings():
        warnings.simplefilter("ignore")
        wb = openpyxl.load_workbook(src)
    n_players = 0
    for row in wb["Roster"].iter_rows(min_row=5, values_only=True):
        if row[0]:
            name = str(row[0]).strip()
            pid = slugify(name)
            if get_player(pid) is None:
                add_player(name, position=row[1] or "RHP")
            update_player(pid, avg_top_velo=row[3], peak_velo=row[2])
            n_players += 1
    colmap = {1: "date", 3: "session_type", 4: "distance", 5: "start_time", 6: "end_time", 8: "actual_throws",
              10: "velo_low", 11: "velo_high", 18: "warm_throws", 19: "warm_velo_low", 20: "warm_velo_high",
              24: "sixty_throws", 25: "sixty_velo_low", 26: "sixty_velo_high", 30: "notes"}
    ws = wb["Throw Log"]
    by_player: dict = {}
    for r in range(5, ws.max_row + 1):
        name = ws.cell(r, 2).value
        if not name:
            continue
        rec = {}
        for col, k in colmap.items():
            v = ws.cell(r, col).value
            rec[k] = None if isinstance(v, str) and v.startswith("=") else v
        rec["created"] = 1000 + r
        by_player.setdefault(slugify(str(name)), []).append(rec)
    n = 0
    for pid, recs in by_player.items():
        if get_player(pid) is None:
            add_player(pid.replace("-", " ").title())
        _run("DELETE FROM throws WHERE player=:p", p=pid)   # replace, never duplicate
        for rec in recs:
            add_row(pid, rec)
            n += 1
    if "Next Outing Planner" in wb.sheetnames:
        nop = wb["Next Outing Planner"]
        if nop["B4"].value and isinstance(nop["B6"].value, datetime):
            update_player(slugify(str(nop["B4"].value)), target_type=nop["B5"].value or "Game",
                          target_date=nop["B6"].value.date(), planned_pitches=nop["B23"].value,
                          planned_innings=nop["D23"].value)
    return n_players, n
